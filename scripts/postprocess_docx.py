import argparse
import json

from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement, ns
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    parser = argparse.ArgumentParser(description="Post-process a DOCX document.")
    parser.add_argument("input", help="Path to the input DOCX file.")
    parser.add_argument("output", help="Path to the output DOCX file.")
    parser.add_argument("--template", help="Template JSON.")

    args = parser.parse_args()

    doc = Document(args.input)

    template = json.loads(args.template) if args.template else {}
    show_page_numbers = template.get("showPageNumbers", True)
    header_text = template.get("headerText", "")
    footer_text = template.get("footerText", "")

    for section in doc.sections:
        header = section.header
        remove_paragraphs(header)
        header_para = header.add_paragraph(header_text)
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        footer = section.footer
        remove_paragraphs(footer)

        if show_page_numbers:
            page_number_para = footer.add_paragraph()
            page_number_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            page_number_para.add_run("Page ")
            add_field(page_number_para, "PAGE")
            page_number_para.add_run(" of ")
            add_field(page_number_para, "NUMPAGES")

        # footer_para = footer.add_paragraph(footer_text)
        # footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    page_width = doc.sections[0].page_width or Inches(8.5)
    left_margin = doc.sections[0].left_margin or Inches(1)
    right_margin = doc.sections[0].right_margin or Inches(1)
    max_width = page_width - left_margin - right_margin

    page_height = doc.sections[0].page_height or Inches(11)
    top_margin = doc.sections[0].top_margin or Inches(1.2)
    bottom_margin = doc.sections[0].bottom_margin or Inches(1)
    max_height = page_height - top_margin - bottom_margin

    for shape in doc.inline_shapes:
        if shape.width > max_width:
            previous_width = shape.width
            shape.width = max_width
            shape.height = int(shape.height * (shape.width / previous_width))
        if shape.height > max_height:
            previous_height = shape.height
            shape.height = max_height
            shape.width = int(shape.width * (shape.height / previous_height))

    doc.save(args.output)

def create_element(name):
    return OxmlElement(name)

def create_attribute(element, name, value):
    element.set(ns.qn(name), value)

def remove_paragraphs(container):
    for para in container.paragraphs:
        p = para._element
        p.getparent().remove(p)
        p._p = p._element = None

def add_field(para, field_code):
    run = para.add_run()

    fldChar1 = create_element('w:fldChar')
    create_attribute(fldChar1, 'w:fldCharType', 'begin')

    instrText = create_element('w:instrText')
    create_attribute(instrText, 'xml:space', 'preserve')
    instrText.text = field_code

    fldChar2 = create_element('w:fldChar')
    create_attribute(fldChar2, 'w:fldCharType', 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

main()

